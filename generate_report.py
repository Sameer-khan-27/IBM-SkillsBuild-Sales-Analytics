import os
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

df = pd.read_csv("SuperMarket Analysis - Copy.csv")
sns.set_theme(style="whitegrid")

# 1. Generate Figures
# Figure 1: Product Line Revenue
plt.figure(figsize=(7, 3.8))
cat_sales = df.groupby('Product line')['Sales'].sum().sort_values(ascending=False)
sns.barplot(x=cat_sales.values, y=cat_sales.index, palette="Blues_r")
plt.title("Total Revenue by Product Line", fontsize=12, weight='bold')
plt.xlabel("Revenue ($)")
plt.ylabel("Product Line")
plt.tight_layout()
plt.savefig("chart_category.png", dpi=200)
plt.close()

# Figure 2: Payment Distribution
plt.figure(figsize=(4.5, 4.5))
df['Payment'].value_counts().plot(kind='pie', autopct='%1.1f%%', colors=sns.color_palette("pastel"), startangle=140)
plt.title("Payment Channel Distribution", fontsize=12, weight='bold')
plt.ylabel("")
plt.tight_layout()
plt.savefig("chart_payment.png", dpi=200)
plt.close()

# Figure 3: Branch Performance
plt.figure(figsize=(6, 3.5))
branch_df = df.groupby('City')['Sales'].sum().reset_index()
sns.barplot(data=branch_df, x='City', y='Sales', palette="viridis")
plt.title("Revenue by Store Location", fontsize=12, weight='bold')
plt.xlabel("City")
plt.ylabel("Revenue ($)")
plt.tight_layout()
plt.savefig("chart_branch.png", dpi=200)
plt.close()

# 2. Build Word Document
doc = Document()
for s in doc.sections:
    s.top_margin = Inches(1)
    s.bottom_margin = Inches(1)
    s.left_margin = Inches(1)
    s.right_margin = Inches(1)

# Title Header
h = doc.add_heading('Supermarket Sales & Business Analytics Report', level=0)
h.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Author Details
meta = doc.add_paragraph()
meta.add_run("Author: ").bold = True
meta.add_run("Md Sameer Khan\n")
meta.add_run("Program: ").bold = True
meta.add_run("AICTE | IBM SkillsBuild Data Analytics with AI Internship\n")
meta.add_run("Institution: ").bold = True
meta.add_run("Maharaja Agrasen Institute of Technology\n")

doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph(
    "This data analytics capstone evaluates 1,000 retail transaction records from supermarket branches across Yangon, "
    "Mandalay, and Naypyitaw. The business objectives are to evaluate category profitability, compare branch operational "
    "performance, examine payment channel preferences, and determine spending behaviors between member and non-member customer segments."
)

doc.add_heading('2. Dataset Overview & Key Metrics', level=1)
doc.add_paragraph(
    "• Total Recorded Revenue: $322,966.75 across 1,000 transactions.\n"
    "• Average Order Value: $322.97 with an average customer rating of 6.97 / 10.0.\n"
    "• Top Performing Product Line: Food and beverages ($56,144.84), followed by Sports and travel ($55,122.83).\n"
    "• Leading Branch: Giza (Naypyitaw) generating $110,568.71 in sales."
)

doc.add_heading('3. Visual Analytics & Findings', level=1)

doc.add_heading('A. Product Line Contribution', level=2)
doc.add_paragraph("Sales distribution across the six core product categories:")
doc.add_picture("chart_category.png", width=Inches(5.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('B. Payment Channel Adoption', level=2)
doc.add_paragraph("Customer checkout preferences reflect strong adoption of modern cashless mechanisms:")
doc.add_picture("chart_payment.png", width=Inches(4.2))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('C. Geographic Branch Comparison', level=2)
doc.add_paragraph("Sales comparison across store branch locations:")
doc.add_picture("chart_branch.png", width=Inches(5.2))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('4. Business Recommendations', level=1)
doc.add_paragraph(
    "1. Stock Replenishment: Prioritize high-turnover lines ('Food and beverages' and 'Sports and travel') to prevent stockout events.\n"
    "2. Loyalty Upselling: Members have a significantly higher basket average ($335.74 vs. $306.37). Launch tiered rewards to convert normal walk-ins.\n"
    "3. Digital Checkout Optimization: With E-wallet (34.5%) and Credit card (31.1%) comprising ~65.6% of volume, offer dedicated self-checkout lanes."
)

doc.add_heading('5. Conclusion', level=1)
doc.add_paragraph(
    "Deploying an interactive dashboard alongside structured data analysis provides management with real-time operational visibility, "
    "ensuring improved inventory management and higher customer retention."
)

out_file = "MdSameerKhan_ProjectReport.docx"
doc.save(out_file)
print(f"Report generated successfully: {out_file}")

# Clean up temporary plots
for f in ["chart_category.png", "chart_payment.png", "chart_branch.png"]:
    if os.path.exists(f):
        os.remove(f)